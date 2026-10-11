import json
import pytest
from openpyxl import Workbook,load_workbook
from openpyxl.worksheet.table import Table,TableStyleInfo
from selected_journal.pdf_sync import sync_pdfs,HEADERS
from selected_journal import cli

def base(tmp_path):
    p=tmp_path/'old.xlsx';wb=Workbook();sh=wb.active
    for c,h in enumerate(HEADERS,1):sh.cell(7,c,h)
    for c,v in enumerate([2020,'Old','Original title','Nature',1.3,0],1):sh.cell(8,c,v)
    t=Table(displayName='Literature',ref='A7:F8');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2');sh.add_table(t)
    sh.freeze_panes='A8';wb.save(p);wb.close();return p

def snapshot(tmp_path,rows):
    p=tmp_path/'snapshot.json';p.write_text(json.dumps({'papers':rows}));return p

def record(key,pdf=True):
    return {'year':2026,'cite_key':key,'title':'=inert title','journal':'Nature','relevance':.8,'citations_per_year':None,'pdf':pdf}

def test_append_pdf_preserves_existing_order_zero_and_native_features(tmp_path):
    p=base(tmp_path);original=p.read_bytes();out=tmp_path/'new.xlsx'
    s=snapshot(tmp_path,[record('Old'),record('New'),record('NoPDF',False),record('Pending','…')])
    report=sync_pdfs(p,s,out)
    assert report['new_keys']==['New'] and report['total_count']==2
    assert p.read_bytes()==original
    wb=load_workbook(out);sh=wb.active
    assert [sh.cell(8,c).value for c in range(1,7)]==[2020,'Old','Original title','Nature',1.3,0]
    assert sh['C9'].data_type=='s' and sh['C9'].value=='=inert title'
    assert sh['F9'].value is None and sh.freeze_panes=='A8'
    assert sh.tables['Literature'].ref==sh.tables['Literature'].autoFilter.ref=='A7:F9'
    wb.close()

def test_duplicate_snapshot_rejected_without_output(tmp_path):
    p=base(tmp_path);s=snapshot(tmp_path,[record('New'),record('New')]);out=tmp_path/'out.xlsx'
    with pytest.raises(ValueError,match='重复'):sync_pdfs(p,s,out)
    assert not out.exists()

def test_absent_pdf_state_and_invalid_metrics_rejected(tmp_path):
    p=base(tmp_path);r=record('New');r.pop('pdf')
    with pytest.raises(ValueError,match='PDF状态'):sync_pdfs(p,snapshot(tmp_path,[r]),tmp_path/'new.xlsx')
    r=record('New');r['year']='2026'
    with pytest.raises(ValueError,match='数值'):sync_pdfs(p,snapshot(tmp_path,[r]),tmp_path/'new.xlsx')

def test_overwrite_disallowed(tmp_path):
    p=base(tmp_path);s=snapshot(tmp_path,[record('New')])
    with pytest.raises(ValueError,match='新xlsx'):sync_pdfs(p,s,p)

def test_supplemental_fields_removed_from_new_query_source_only(tmp_path):
    from selected_journal.enrichment import EXPORT_LABELS
    from openpyxl.worksheet.table import TableColumn
    p=base(tmp_path);wb=load_workbook(p);sh=wb.active;t=sh.tables['Literature']
    for c,label in enumerate(EXPORT_LABELS,7):sh.cell(7,c,label);sh.cell(8,c,'old metadata');t.tableColumns.append(TableColumn(id=c,name=label))
    t.ref='A7:BP8';t.autoFilter.ref=t.ref;sh.merge_cells('G2:Q2');sh['G2']='Old metadata note';wb.save(p);wb.close()
    original=p.read_bytes();out=tmp_path/'new.xlsx'
    assert cli.main(['sync-pdfs','--input',str(p),'--snapshot',str(snapshot(tmp_path,[record('New')])),'--output',str(out)])==0
    wb=load_workbook(out);sh=wb.active
    assert sh.max_column==6 and len(sh.tables['Literature'].tableColumns)==6
    assert not sh.merged_cells.ranges and sh['B8'].value=='Old' and sh['B9'].value=='New'
    assert p.read_bytes()==original;wb.close()
