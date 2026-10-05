import copy
import csv
import http.cookiejar
import json
from pathlib import Path

import pytest
import requests
from openpyxl import Workbook, load_workbook

from selected_journal import cli
from selected_journal.client import AccessStopped, FetchError, LetPubClient, load_cookies
from selected_journal.inputs import make_row, normalize_issn, read_file, read_lines
from selected_journal.parsing import ParseError, parse_detail, parse_search
from selected_journal.runner import resolve, run_batch
from selected_journal.storage import read_events

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def journal():
    return parse_detail((FIXTURES / "detail_anonymous.html").read_text(), "6969")


def test_real_anonymous_detail_and_hidden_partitions(journal):
    f = journal["fields"]
    assert f["journal_name"]["value"] == "PROTEIN AND PEPTIDE LETTERS"
    assert f["issn"]["value"] == "0929-8665"
    assert f["eissn"]["value"] == "1875-5305"
    assert f["citescore"]["value"] == 2.8
    assert f["impact_factor"]["value"] is None
    assert f["impact_factor"]["status"] == "login_required"
    assert f["citescore"]["year"] is None  # release date isn't metric year
    assert f["review_time"]["value"] == "网友分享经验： 平均2.0个月"
    assert f["oa"]["value"] == "closed"
    p = {p["scheme"]: p for p in journal["classifications"]}
    assert p["jcr"]["status"] == "login_required"
    assert [e["quartile"] for e in p["xinrui"]["entries"]] == ["4区", "4区"]
    assert [e["quartile"] for e in p["journal_partition"]["entries"]] == ["4区", "4区"]
    assert p["journal_partition"]["year"] == 2025


def detail_with(extra):
    return ('<table><tr><td>期刊名字</td><td><a href="?page=journalapp&view=detail&journalid=6969">PROTEIN AND PEPTIDE LETTERS</a></td></tr>'
            '<tr><td>期刊ISSN</td><td>0929-8665</td></tr>' + extra + '</table>')


def test_logged_in_metrics_multisubject_and_css_visibility():
    html = detail_with('''
    <tr><td>2025年度IF</td><td>3.5</td></tr>
    <tr><td>实时影响因子</td><td>999.9</td></tr>
    <tr><td>WOS期刊JCR分区（2026最新版）</td><td>
    <table><tr><th>按JIF指标学科分区</th><th>分区</th></tr><tr><td>Biology</td><td><span class="decoy">Q4</span>Q1</td></tr></table>
    <table><tr><th>按JCI指标学科分区</th><th>分区</th></tr><tr><td>Biology</td><td>Q2</td></tr><tr><td>Chemistry</td><td>Q3</td></tr></table>
    </td></tr><style>.decoy {display:none}</style>''')
    # The explicit metric year is retained; real-time IF is ignored.
    r = parse_detail(html, "6969")
    assert r["fields"]["impact_factor"]["value"] == 3.5
    assert r["fields"]["impact_factor"]["year"] == 2025
    jcr = next(p for p in r["classifications"] if p["scheme"] == "jcr")
    assert [(e["basis"], e["subject"], e["quartile"]) for e in jcr["entries"]] == [
        ("按JIF指标学科分区", "Biology", "Q1"), ("按JCI指标学科分区", "Biology", "Q2"), ("按JCI指标学科分区", "Chemistry", "Q3")]
    assert r["fields"]["review_time"]["status"] == "not_provided"


def test_invalid_detail_field_not_filled_from_other_numbers():
    r = parse_detail(detail_with('<tr><td>2026最新IF</td><td>错误页 123</td></tr>'), "6969")
    assert r["fields"]["impact_factor"]["value"] is None
    assert r["fields"]["impact_factor"]["status"] == "parse_error"
    with pytest.raises(ParseError):
        parse_detail('<h1>Website updated</h1>', "6969")


def test_real_search_excludes_recommendations_articles_and_last_page_link():
    search = parse_search((FIXTURES / "search.html").read_text())
    assert [c["journal_id"] for c in search["candidates"]] == ["6969"]
    assert search["next_url"] is None


def search_html(names=(("6969", "PROTEIN AND PEPTIDE LETTERS"),), page=None):
    links = ''.join(f'<tr><td>0929-8665</td><td><a href="?page=journalapp&view=detail&journalid={i}">{n}</a></td></tr>' for i, n in names)
    if page:
        links += f'<tr><td><a href="?page=journalapp&view=search&currentsearchpage={page}">下一页</a></td></tr>'
    return '<table class="table_yjfx"><tr><th>ISSN</th><th>期刊名</th></tr>'+links+'</table>'


def test_search_pagination_and_distinguish_empty_from_bad_page():
    assert "currentsearchpage=2" in parse_search(search_html(page=2))["next_url"]
    assert parse_search(search_html(names=()))["candidates"] == []
    assert parse_search('<p>未找到期刊记录</p>')["candidates"] == []
    with pytest.raises(ParseError):
        parse_search('<h1>Gateway error</h1>')


@pytest.mark.parametrize("raw,expected", [("09298665", "0929-8665"), ("0000-0000", "0000-0000"), ("１２３４-５６７９", "1234-5679"), ("2434561x", "2434-561X")])
def test_issn_formats(raw, expected):
    assert normalize_issn(raw) == expected


def test_bad_issn_stays_invalid():
    rows = read_lines(['  Nature  ', '', '0929-8666', '0929-8665'])
    assert len(rows) == 3
    assert rows[1]["error"]
    assert rows[1]["journal_name"] == ''
    assert rows[2]["issn"] == '0929-8665'
    with pytest.raises(ValueError):
        normalize_issn('bad')


def test_file_inputs_and_first_sheet(tmp_path):
    (tmp_path/'j.txt').write_text('\ufeffNature\n0929-8665\n')
    (tmp_path/'j.csv').write_text('journal_name,issn\nNature,\n,0929-8665\n,\n')
    wb = Workbook()
    wb.active.append(['journal_name','issn'])
    wb.active.append(['Nature',None])
    wb.active.append([None,'0929-8665'])
    wb.create_sheet('ignored').append(['not valid columns'])
    wb.save(tmp_path/'j.xlsx')
    for suffix in ('txt','csv','xlsx'):
        rows = read_file(tmp_path/f'j.{suffix}')
        assert [r['journal_name'] for r in rows] == ['Nature','']
        assert rows[1]['issn'] == '0929-8665'
    (tmp_path/'bad.csv').write_text('title\nNature\n')
    with pytest.raises(ValueError):
        read_file(tmp_path/'bad.csv')


class FakeClient:
    def __init__(self, journal):
        self.record = journal
        self.search_calls = 0
        self.detail_calls = 0
        self.candidates = [{'journal_id':'6969','journal_name':'PROTEIN AND PEPTIDE LETTERS','issn':'0929-8665','url':'https://www.letpub.com.cn/'}]
        self.complete = True

    def search(self, row):
        self.search_calls += 1
        return {'candidates':self.candidates,'complete':self.complete}

    def detail(self, identifier):
        self.detail_calls += 1
        return copy.deepcopy(self.record)


def test_matching_conflicts_ambiguity_eissn_and_empty(journal):
    c = FakeClient(journal)
    assert resolve(make_row(1,issn='1875-5305'), c, {})['query']['status'] == 'matched'
    assert resolve(make_row(2,'Wrong Journal','0929-8665'), c, {})['query']['status'] == 'needs_confirmation'
    assert resolve(make_row(3,'Protein'), c, {})['candidates']
    c.complete = False
    assert resolve(make_row(4,issn='0929-8665'), c, {})['query']['status'] == 'needs_confirmation'
    c.complete = True
    c.candidates = []
    assert resolve(make_row(5,'Unknown'), c, {})['query']['status'] == 'not_found'


def test_dedup_resume_and_exports_agree(tmp_path, journal):
    rows = [make_row(1,'PROTEIN AND PEPTIDE LETTERS'), make_row(2,issn='0929-8665'), make_row(3,issn='1875-5305')]
    c = FakeClient(journal)
    code, data = run_batch(rows, tmp_path, c, emit=lambda _:None)
    assert code == 0 and len(data['journals']) == 1
    assert data['journals'][0]['input_ids'] == [1,2,3]
    assert c.search_calls == c.detail_calls == 1
    run_batch(rows, tmp_path, c, read_events(tmp_path), emit=lambda _:None)
    assert c.search_calls == 1
    assert json.loads((tmp_path/'results.json').read_text()) == data
    wb = load_workbook(tmp_path/'results.xlsx',read_only=True)
    assert wb.sheetnames == ['期刊汇总','分区明细','查询记录','候选清单']
    for sheet in wb:
        csv_file = tmp_path/f'{sheet.title}.csv'
        assert csv_file.read_bytes().startswith(b'\xef\xbb\xbf')
        with csv_file.open(encoding='utf-8-sig',newline='') as stream:
            csv_rows = list(csv.reader(stream))
        xlsx_rows = [[str(v) if v is not None else '' for v in row] for row in sheet.values]
        assert csv_rows == xlsx_rows
    wb.close()


@pytest.mark.parametrize('exception,status,code', [(AccessStopped('captcha'),'blocked',2),(KeyboardInterrupt(),'pending',130),(FetchError('timeout'),'failed',2)])
def test_failure_interrupt_and_checkpoint(tmp_path,journal,exception,status,code):
    rows = [make_row(1,'PROTEIN AND PEPTIDE LETTERS'),make_row(2,'Other'),make_row(3,'Third')]
    c = FakeClient(journal)
    original = c.search
    def search(row):
        if row['input_id']>=2: raise exception
        return original(row)
    c.search = search
    result_code,data = run_batch(rows,tmp_path,c,emit=lambda _:None)
    assert result_code == code
    assert data['queries'][0]['status'] == 'matched'
    assert data['queries'][1]['status'] == status
    assert len(data['journals']) == 1
    events = read_events(tmp_path)
    assert events[1]['journal']['journal_id'] == '6969'
    assert (tmp_path/'results.xlsx').exists()
    if status=='blocked': assert data['queries'][2]['status'] == 'pending'


def test_trailing_checkpoint_recovery(tmp_path,journal):
    run_batch([make_row(1,'PROTEIN AND PEPTIDE LETTERS')],tmp_path,FakeClient(journal),emit=lambda _:None)
    with (tmp_path/'checkpoint.jsonl').open('ab') as f: f.write(b'{"query":')
    assert list(read_events(tmp_path)) == [1]
    assert (tmp_path/'checkpoint.jsonl').read_bytes().endswith(b'\n')


class Clock:
    def __init__(self): self.time=0; self.sleeps=[]
    def now(self): return self.time
    def sleep(self,seconds): self.sleeps.append(seconds); self.time+=seconds


def fake_session(outcomes):
    session = requests.Session()
    calls = []
    def request(method,url,**kwargs):
        calls.append((method,url,kwargs))
        outcome = outcomes.pop(0)
        if isinstance(outcome,Exception): raise outcome
        response = requests.Response()
        response.status_code,response._content = outcome[0],outcome[1].encode()
        return response
    session.request = request
    return session,calls


def test_http_retry_timeout_spacing_and_nontransient(tmp_path):
    clock=Clock()
    session,calls=fake_session([requests.Timeout(),(503,'oops'),(200,'ok'),(200,'ok'),(404,'missing')])
    c=LetPubClient(tmp_path,session=session,sleep=clock.sleep,clock=clock.now)
    assert c.fetch('GET','https://www.letpub.com.cn/')=='ok'
    assert 5 in clock.sleeps and 15 in clock.sleeps
    c.fetch('GET','https://www.letpub.com.cn/')
    assert clock.sleeps[-1]==3
    assert calls[0][2]['timeout']==(10,30)
    with pytest.raises(FetchError): c.fetch('GET','https://www.letpub.com.cn/')
    assert len(calls)==5


@pytest.mark.parametrize('outcome',[(403,'blocked'),(429,'slow'),(200,'<h1>请输入验证码</h1>')])
def test_access_stop_makes_no_further_requests(tmp_path,outcome):
    session,calls=fake_session([outcome])
    c=LetPubClient(tmp_path,session=session,sleep=lambda _:None)
    for _ in range(2):
        with pytest.raises(AccessStopped): c.fetch('GET','https://www.letpub.com.cn/')
    assert len(calls)==1


def test_cache_refresh_expiry_and_memory_dedup(tmp_path):
    html=(FIXTURES/'detail_anonymous.html').read_text()
    session,calls=fake_session([(200,html)])
    c=LetPubClient(tmp_path,session=session,sleep=lambda _:None)
    first=c.detail('6969')
    first['fields']['issn']['value']='mutated'
    assert c.detail('6969')['fields']['issn']['value']=='0929-8665'
    assert len(calls)==1
    other,calls2=fake_session([])
    assert LetPubClient(tmp_path,session=other).detail('6969')['journal_id']=='6969'
    assert calls2==[]
    fresh,calls3=fake_session([(200,html)])
    LetPubClient(tmp_path,refresh=True,session=fresh).detail('6969')
    assert len(calls3)==1
    cache=tmp_path/'anonymous/6969.json'
    saved=json.loads(cache.read_text()); saved['saved_at']-=86401; cache.write_text(json.dumps(saved))
    expired,calls4=fake_session([(200,html)])
    LetPubClient(tmp_path,session=expired).detail('6969')
    assert len(calls4)==1


def test_paginated_search_and_incomplete_cap(tmp_path):
    pages=[(200,'search form'),(200,search_html(page=2)),(200,search_html(names=(('123','SECOND'),)))]
    s,calls=fake_session(pages)
    c=LetPubClient(tmp_path,session=s,sleep=lambda _:None)
    r=c.search(make_row(1,'PROTEIN AND PEPTIDE LETTERS'))
    assert r['complete'] and len(r['candidates'])==2
    assert calls[1][2]['data']['searchname']=='PROTEIN AND PEPTIDE LETTERS'
    c.search(make_row(2,'PROTEIN AND PEPTIDE LETTERS'))
    assert len(calls)==3
    pages=[(200,'form')]+[(200,search_html(names=((str(i),'NAME'),),page=i+1)) for i in range(1,11)]
    s,calls=fake_session(pages)
    c=LetPubClient(tmp_path,session=s,sleep=lambda _:None)
    assert not c.search(make_row(1,'NAME'))['complete']
    assert len(calls)==11


def test_netscape_cookie_filter_and_cache_isolation(tmp_path):
    cookie=tmp_path/'local.txt'
    cookie.write_text('# Netscape HTTP Cookie File\n.letpub.com.cn\tTRUE\t/\tTRUE\t2147483647\tsid\tSECRET\n.evil.test\tTRUE\t/\tFALSE\t2147483647\tx\tOTHER\n')
    session=requests.Session()
    mode=load_cookies(session,cookie)
    assert mode.startswith('login-') and 'SECRET' not in mode
    assert [(c.name,c.value) for c in session.cookies]==[('sid','SECRET')]
    html=(FIXTURES/'detail_anonymous.html').read_text()
    session,calls=fake_session([(200,html)])
    client=LetPubClient(tmp_path,cookie_file=cookie,session=session)
    client.detail('6969')
    assert client.cache_dir.name==mode
    assert 'SECRET' not in next(client.cache_dir.glob('*.json')).read_text()
    assert not (tmp_path/'anonymous').exists()
    bad=tmp_path/'bad.txt'; bad.write_text('incorrect')
    with pytest.raises(ValueError): load_cookies(requests.Session(),bad)


def test_cli_errors_and_stdin(tmp_path,monkeypatch,journal):
    assert cli.main(['query'])==1
    assert cli.main(['query','--stdin','--resume',str(tmp_path)])==1
    assert cli.main(['query','--resume',str(tmp_path),'--output',str(tmp_path)])==1
    class FakeCLIClient(FakeClient):
        mode='anonymous'
        def __init__(self,*a,**kw): super().__init__(journal)
        def close(self): pass
    import io
    monkeypatch.setattr(cli,'LetPubClient',FakeCLIClient)
    monkeypatch.setattr('sys.stdin',io.StringIO(''))
    assert cli.main(['query','--stdin','--output',str(tmp_path/'empty')])==1
    assert not (tmp_path/'empty').exists()
    monkeypatch.setattr('sys.stdin',io.StringIO('0929-8665\n'))
    assert cli.main(['query','--stdin','--output',str(tmp_path/'run')])==0
    assert cli.main(['query','--resume',str(tmp_path/'run')])==0
    monkeypatch.setattr('sys.stdin',io.StringIO('0929-8665\n'))
    assert cli.main(['query','--stdin','--output',str(tmp_path/'run')])==1


def test_latest_partition_selected_even_if_old_block_comes_first():
    def block(year,q):
        return f'<tr><td>期刊分区表（{year}年3月升级版）</td><td><table><tr><th>大类学科</th><th>小类学科</th></tr><tr><td>Biology {q}区</td><td><table><tr><td>Cell</td><td>{q}区</td></tr></table></td></tr></table></td></tr>'
    r=parse_detail(detail_with(block(2023,4)+block(2025,2)), '6969')
    p=next(p for p in r['classifications'] if p['scheme']=='journal_partition')
    assert p['year']==2025
    assert [e['quartile'] for e in p['entries']]==['2区','2区']


def test_failed_and_partial_inputs_retry_on_resume(tmp_path,journal):
    rows=[make_row(1,'PROTEIN AND PEPTIDE LETTERS')]
    bad=copy.deepcopy(journal)
    bad['fields']['citescore']['status']='parse_error'
    c=FakeClient(bad)
    assert run_batch(rows,tmp_path,c,emit=lambda _:None)[0]==2
    recovered=FakeClient(journal)
    code,data=run_batch(rows,tmp_path,recovered,read_events(tmp_path),emit=lambda _:None)
    assert code==0 and data['queries'][0]['status']=='matched'
    assert recovered.detail_calls==1
    assert data['journals'][0]['fields']['citescore']['status']=='ok'


def test_malformed_xlsx_and_ambiguous_full_name(tmp_path,journal):
    bad=tmp_path/'bad.xlsx';bad.write_text('not an xlsx')
    with pytest.raises(ValueError):read_file(bad)
    c=FakeClient(journal)
    c.candidates=c.candidates+[dict(c.candidates[0],journal_id='123')]
    e=resolve(make_row(1,'PROTEIN AND PEPTIDE LETTERS'),c,{})
    assert e['query']['status']=='needs_confirmation'
    assert len(e['candidates'])==2
    assert c.detail_calls==0


def test_invalid_cookie_can_load_public_information(tmp_path):
    path=tmp_path/'expired.txt'
    path.write_text('# Netscape HTTP Cookie File\n.letpub.com.cn\tTRUE\t/\tTRUE\t1\tsid\tEXPIRED\n')
    assert load_cookies(requests.Session(),path)=='anonymous'


def test_repeated_pagination_and_failed_later_page_need_confirmation(tmp_path):
    page=search_html(page=2)
    for tail in [(200,page),(404,'not a page')]:
        s,calls=fake_session([(200,'form'),(200,page),tail])
        client=LetPubClient(tmp_path,session=s,sleep=lambda _:None)
        result=client.search(make_row(1,'PROTEIN AND PEPTIDE LETTERS'))
        assert not result['complete'] and len(result['candidates'])==1


def test_transient_error_attempts_are_bounded(tmp_path):
    s,calls=fake_session([requests.Timeout(),requests.Timeout(),requests.Timeout()])
    client=LetPubClient(tmp_path,session=s,sleep=lambda _:None)
    with pytest.raises(FetchError):client.fetch('GET','https://www.letpub.com.cn/')
    assert len(calls)==3


def test_partial_parse_is_not_persisted_in_success_cache(tmp_path):
    html=detail_with('<tr><td>2026最新IF</td><td>Unexpected response</td></tr>')
    s,calls=fake_session([(200,html)])
    client=LetPubClient(tmp_path,session=s)
    assert client.detail('6969')['fields']['impact_factor']['status']=='parse_error'
    assert not list(tmp_path.rglob('*.json'))


def test_exported_spreadsheet_text_is_inert_and_json_preserves_input(tmp_path,journal):
    from selected_journal.storage import export_results, result_data
    row=make_row(1,'=1+1',raw='=1+1')
    data=result_data([row],{})
    export_results(tmp_path,data)
    wb=load_workbook(tmp_path/'results.xlsx')
    value=wb['查询记录']['B2']
    assert value.data_type=='s' and value.value=="'=1+1"
    wb.close()
    assert json.loads((tmp_path/'results.json').read_text())['queries'][0]['journal_name']=='=1+1'
