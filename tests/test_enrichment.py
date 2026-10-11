import json
from pathlib import Path
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from selected_journal import cli
from selected_journal.enrichment import extract, prepare, join_rows, EXPORT_KEYS, EXPORT_LABELS
from selected_journal.parsing import parse_detail
from selected_journal.storage import result_data
from selected_journal.inputs import make_row
from selected_journal.runner import make_event


def literature(path, entries, header=7):
    wb=Workbook()
    sh=wb.active
    sh.title="文献"
    sh.cell(2,1,"研究文献")
    for col,value in enumerate(["Cite key","所属期刊","备注"],1): sh.cell(header,col,value)
    for row,entry in enumerate(entries,header+1):
        for col,value in enumerate(entry,1): sh.cell(row,col,value)
    sh.freeze_panes=f"A{header+1}"
    sh.column_dimensions["A"].width=18
    table=Table(displayName="Literature",ref=f"A{header}:C{header+len(entries)}")
    table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
    sh.add_table(table)
    wb.save(path)
    return path


@pytest.fixture
def results(tmp_path):
    rec=parse_detail((Path(__file__).parent/"fixtures/detail_anonymous.html").read_text(),"6969")
    rec["fields"]["impact_factor"]={"value":0,"status":"ok","raw":"0","version":"2025-2026最新IF","year":None}
    rows=[make_row(1,"PROTEIN AND PEPTIDE LETTERS")]
    data=result_data(rows,{1:make_event(rows[0],"matched","精确匹配",rec)})
    path=tmp_path/"results.json"
    path.write_text(json.dumps(data))
    return path


def test_dedup_case_space_nfkc_and_skip_platforms(tmp_path):
    p=literature(tmp_path/"in.xlsx",[["A","Protein And Peptide Letters",None],
        ["B","  PROTEIN AND PEPTIDE LETTERS  ",None],["C","bioRxiv",None],
        ["D","Unknown Journal",None],["E","ＰＲＯＴＥＩＮ AND PEPTIDE LETTERS",None]])
    wb,sh,m=extract(p)
    stats=prepare(m,tmp_path/"work")
    assert stats["papers"]==5 and stats["queries"]==1
    assert m["header_row"]==7
    assert [r["record_type"] for r in m["rows"]][-3:]==["preprint","unknown","journal"]
    wb.close()


def test_duplicate_keys_and_formula_names_rejected(tmp_path):
    p=literature(tmp_path/"dup.xlsx",[["A","Nature",None],["A","Cell",None]])
    with pytest.raises(ValueError,match="重复"): extract(p)
    p=literature(tmp_path/"formula.xlsx",[["A","=B1",None]])
    with pytest.raises(ValueError,match="公式"): extract(p)


def test_invalid_mapping_issn_and_missing_provenance(tmp_path):
    p=tmp_path/"m.json"
    p.write_text(json.dumps({"verified_aliases":{"Test":{"name":"Nature","issn":"0000-0001","source":"https://example.com"}}}))
    with pytest.raises(ValueError,match="校验"): extract(literature(tmp_path/"a.xlsx",[["A","Test",None]]),mappings=p)
    p.write_text(json.dumps({"verified_aliases":{"Test":{"name":"Nature"}}}))
    with pytest.raises(ValueError,match="source"): extract(tmp_path/"a.xlsx",mappings=p)


def test_join_zero_missing_version_no_guess_and_preprint(tmp_path,results):
    p=literature(tmp_path/"in.xlsx",[["A","PROTEIN AND PEPTIDE LETTERS",None],["B","bioRxiv",None],
                                 ["C","Unresolved journal",None]])
    wb,sh,m=extract(p)
    values,mapping=join_rows(m,[results])
    assert values[0][EXPORT_KEYS.index("impact_factor")]==0
    assert values[0][EXPORT_KEYS.index("impact_factor_year")] is None
    assert "2025" in values[0][EXPORT_KEYS.index("impact_factor_version")]
    assert mapping[1]["match_status"]=="preprint"
    assert values[1][EXPORT_KEYS.index("impact_factor")] is None
    assert mapping[2]["match_status"]=="pending"
    wb.close()


def test_offline_cli_preserves_original_table_styles_freeze_and_other_sheet(tmp_path,results):
    p=literature(tmp_path/"in.xlsx",[["A","PROTEIN AND PEPTIDE LETTERS","keep"]])
    wb=load_workbook(p);wb.create_sheet("Other")["A1"]="unchanged";wb.save(p);wb.close()
    original=p.read_bytes()
    out=tmp_path/"v2.xlsx"
    assert cli.main(["enrich","--input",str(p),"--output",str(out),"--results",str(results)])==0
    assert p.read_bytes()==original
    wb=load_workbook(out)
    sh=wb["文献"]
    assert wb.sheetnames==["文献","Other"] and wb["Other"]["A1"].value=="unchanged"
    assert [sh.cell(8,c).value for c in range(1,4)]==["A","PROTEIN AND PEPTIDE LETTERS","keep"]
    assert sh.freeze_panes=="A8" and sh.column_dimensions["A"].width==18
    table=sh.tables["Literature"]
    assert table.tableStyleInfo.name=="TableStyleMedium2"
    assert len(table.tableColumns)==3+len(EXPORT_KEYS)
    assert len({c.name for c in table.tableColumns})==len(table.tableColumns)
    assert table.autoFilter.ref==table.ref
    expected=json.loads((tmp_path/"v2_资料"/"enrichment_values.json").read_text())["values"][0]
    assert [sh.cell(8,c).value for c in range(4,4+len(EXPORT_KEYS))]==expected
    wb.close()
    assert cli.main(["enrich","--input",str(p),"--output",str(p),"--results",str(results)])==1
    assert cli.main(["enrich","--input",str(p),"--output",str(out),"--results",str(results)])==1


def test_verified_alias_and_paper_override_resolve_via_issn(tmp_path,results):
    p=literature(tmp_path/"in.xlsx",[["A","Long wrong spelling",None],["B","Unknown Journal",None]])
    maps=tmp_path/"mappings.json"
    maps.write_text(json.dumps({"verified_aliases":{"Long wrong spelling":{"name":"PROTEIN AND PEPTIDE LETTERS",
        "issn":"0929-8665","source":"https://publisher.example"}},
        "paper_overrides":{"B":{"standard_name":"PROTEIN AND PEPTIDE LETTERS","issn":"0929-8665",
                               "sources":["https://doi.org/example"],"record_type":"journal"}}}))
    wb,sh,m=extract(p,mappings=maps)
    assert len(m["queries"])==1 and m["queries"][0]["journal_name"]==""
    values,mapping=join_rows(m,[results])
    assert all(row["match_status"]=="matched" for row in mapping)
    assert values[0][EXPORT_KEYS.index("issn")]=="0929-8665"
    assert values[1][EXPORT_KEYS.index("verification_source")]=="https://doi.org/example"
    wb.close()


def test_online_enrich_calls_existing_query_command(tmp_path,results,monkeypatch):
    p=literature(tmp_path/"in.xlsx",[["A","PROTEIN AND PEPTIDE LETTERS",None]])
    called=[]
    def fake_query(inputs,directory,client,events,**kwargs):
        data=json.loads(results.read_text())
        directory.mkdir(parents=True,exist_ok=True)
        (directory/"results.json").write_text(json.dumps(data))
        called.extend(inputs)
        return 0,data
    class FakeClient:
        mode="anonymous"
        def __init__(self,*args,**kwargs): pass
        def close(self): pass
    monkeypatch.setattr(cli,"LetPubClient",FakeClient)
    monkeypatch.setattr(cli,"run_batch",fake_query)
    assert cli.main(["enrich","--input",str(p),"--output",str(tmp_path/"out.xlsx")])==0
    assert len(called)==1


def test_multiple_results_do_not_replace_success_with_failure(tmp_path,results):
    failed=tmp_path/"failed.json"
    data=json.loads(results.read_text());data["journals"]=[]
    data["queries"][0].update(status="failed",journal_id=None,message="network")
    failed.write_text(json.dumps(data))
    p=literature(tmp_path/"in.xlsx",[["A","PROTEIN AND PEPTIDE LETTERS",None]])
    wb,sh,m=extract(p);_,mapping=join_rows(m,[results,failed])
    assert mapping[0]["match_status"]=="matched"
    wb.close()


def test_multi_subject_jif_jci_arrays_preserved(tmp_path,results):
    data=json.loads(results.read_text())
    partition=data["journals"][0]["classifications"][0]
    partition.update(scheme="jcr",status="ok",version="2025-2026",year=None,
        entries=[{"level":"subject","subject":"A","quartile":"Q1","basis":"按JIF学科"},
                 {"level":"subject","subject":"B","quartile":"Q2","basis":"按JCI学科"}])
    results.write_text(json.dumps(data))
    p=literature(tmp_path/"in.xlsx",[["A","PROTEIN AND PEPTIDE LETTERS",None]])
    wb,sh,m=extract(p);values,_=join_rows(m,[results])
    text=values[0][EXPORT_KEYS.index("jcr_subjects")]
    assert "按JIF学科: A Q1" in text and "按JCI学科: B Q2" in text
    assert values[0][EXPORT_KEYS.index("jcr_year")] is None
    wb.close()



def test_numeric_zero_key_is_retained(tmp_path):
    p=literature(tmp_path/"in.xlsx",[[0,"Nature",None]])
    wb,sh,m=extract(p)
    assert m["rows"][0]["key"]=="0"
    wb.close()


def test_already_enriched_rejected_before_query(tmp_path,monkeypatch):
    p=literature(tmp_path/"in.xlsx",[["A","Nature",None]])
    wb=load_workbook(p);wb.active.cell(7,4,EXPORT_LABELS[0]);wb.save(p);wb.close()
    def unexpected(*args,**kwargs):
        raise AssertionError("should validate before network access")
    monkeypatch.setattr(cli,"LetPubClient",unexpected)
    assert cli.main(["enrich","--input",str(p),"--output",str(tmp_path/"out.xlsx")])==1
    assert not (tmp_path/"out.xlsx").exists()


def test_complete_issn_recheck_clears_old_partial_status(tmp_path,results):
    complete=json.loads(results.read_text())
    partial=json.loads(results.read_text())
    partial["queries"][0].update(status="partial",message="部分字段解析失败: oa")
    partial["journals"][0]["fields"]["oa"]["status"]="parse_error"
    initial=tmp_path/"partial.json";initial.write_text(json.dumps(partial))
    complete["queries"][0].update(journal_name="",issn="0929-8665")
    results.write_text(json.dumps(complete))
    p=literature(tmp_path/"in.xlsx",[["A","PROTEIN AND PEPTIDE LETTERS",None]])
    wb,sh,m=extract(p);values,mapping=join_rows(m,[initial,results])
    assert mapping[0]["match_status"]=="matched"
    assert values[0][EXPORT_KEYS.index("oa_status")]=="ok"
    wb.close()
