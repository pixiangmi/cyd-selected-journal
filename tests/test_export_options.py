import copy
import csv
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from selected_journal import cli
from selected_journal.client import atomic_json
from selected_journal.inputs import make_row
from selected_journal.parsing import parse_detail
from selected_journal.runner import make_event, run_batch
from selected_journal.storage import (
    ExportOptions, HEADER_LABELS, append_event, export_results, is_missing,
    missing_statistics, projections, result_data,
)


@pytest.fixture
def data():
    journal = parse_detail((Path(__file__).parent / "fixtures/detail_anonymous.html").read_text(), "6969")
    rows = [make_row(1, issn="0929-8665"), make_row(2, "PROTEIN AND PEPTIDE LETTERS")]
    events = {row["input_id"]: make_event(row, "matched", "精确匹配", journal) for row in rows}
    return result_data(rows, events)


def test_labels_cover_all_tables_without_collisions(data):
    for _, headers, _ in projections(data):
        labels = [HEADER_LABELS[h] for h in headers]
        assert len(labels) == len(set(labels))


@pytest.mark.parametrize("value", [None, "", " \t\n", [], {}, ()])
def test_missing_definition(value):
    assert is_missing(value)


@pytest.mark.parametrize("value", [0, 0.0, False, "unknown", "0", "False", [None]])
def test_falsy_and_unknown_values_are_not_missing(value):
    assert not is_missing(value)


@pytest.mark.parametrize("language", ["original", "zh"])
@pytest.mark.parametrize("drop", [False, True])
@pytest.mark.parametrize("report_mode", ["none", "terminal", "file", "both"])
def test_export_matrix_consistent_tables_and_reports(tmp_path, data, language, drop, report_mode):
    original = copy.deepcopy(data)
    options = ExportOptions(language, drop, report_mode)
    messages = []
    export_results(tmp_path, data, options, emit=messages.append)
    saved = json.loads((tmp_path / "results.json").read_text())
    assert {key: saved[key] for key in data} == data == original
    workbook = load_workbook(tmp_path / "results.xlsx", read_only=True)
    for title, headers, rows in projections(data):
        kept = [h for h in headers if not (drop and rows and all(is_missing(row.get(h)) for row in rows))]
        expected = [options.label(h) for h in kept]
        with (tmp_path / f"{title}.csv").open(encoding="utf-8-sig", newline="") as stream:
            csv_rows = list(csv.reader(stream))
        excel_rows = [[str(v) if v is not None else "" for v in row] for row in workbook[title].values]
        assert csv_rows == excel_rows
        assert csv_rows[0] == expected
        assert len(csv_rows) == len(rows) + 1
        if language == "zh" or drop:
            table = next(t for t in saved["export_view"]["tables"] if t["table"] == title)
            assert table["columns"] == [{"field": h, "label": options.label(h)} for h in kept]
            assert table["rows"] == [{options.label(h): row.get(h) for h in kept} for row in rows]
        else:
            assert "export_view" not in saved
    workbook.close()
    file_report = report_mode in ("file", "both")
    assert (tmp_path / "missing_values.json").exists() == file_report
    assert (tmp_path / "missing_values.csv").exists() == file_report
    assert bool(messages) == (report_mode in ("terminal", "both"))
    if file_report:
        report = json.loads((tmp_path / "missing_values.json").read_text())
        assert report == missing_statistics(data, options)
        with (tmp_path / "missing_values.csv").open(encoding="utf-8-sig", newline="") as stream:
            records = list(csv.reader(stream))
        assert len(records) == sum(len(table["columns"]) for table in report["tables"]) + 1
        assert records[0][0] == ("工作表" if language == "zh" else "table")
        assert (tmp_path / "missing_values.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    assert not list(tmp_path.glob(".tmp-*"))


def test_statistics_include_removed_columns_and_missing_reasons(data):
    # Two input rows point at a single journal; counts use deduplicated table rows.
    report = missing_statistics(data, ExportOptions("zh", True, "file"))
    summary = report["tables"][0]
    assert summary["row_count"] == 1
    columns = {c["field"]: c for c in summary["columns"]}
    assert columns["impact_factor"]["missing_count"] == 1
    assert columns["impact_factor"]["missing_rate"] == 1
    assert columns["impact_factor"]["dropped"]
    assert columns["impact_factor"]["reasons"] == {"login_required": 1}
    assert not columns["impact_factor_status"]["dropped"]
    assert columns["journal_name"]["missing_count"] == 0
    assert report["tables"][2]["row_count"] == 2
    candidates = report["tables"][3]
    assert candidates["row_count"] == 0
    assert all(c["missing_rate"] is None and not c["dropped"] for c in candidates["columns"])


def test_counts_mix_blank_and_zero_and_false(data):
    rows = [{"input_id": i, "journal_name": value} for i, value in enumerate([None, " ", "Nature", 0, False])]
    data["queries"] = rows
    column = next(c for c in missing_statistics(data, ExportOptions())["tables"][2]["columns"]
                  if c["field"] == "journal_name")
    assert column["missing_count"] == 2 and column["present_count"] == 3
    assert column["missing_rate"] == 0.4


def test_empty_task_keeps_schema_and_no_false_missing_rate(tmp_path):
    data = {"schema_version": 1, "journals": [], "queries": [], "candidates": []}
    export_results(tmp_path, data, ExportOptions("zh", True, "both"), emit=lambda _: None)
    wb = load_workbook(tmp_path / "results.xlsx", read_only=True)
    for title, headers, _ in projections(data):
        assert list(next(wb[title].values)) == [HEADER_LABELS[h] for h in headers]
    wb.close()
    report = json.loads((tmp_path / "missing_values.json").read_text())
    assert all(c["missing_rate"] is None for t in report["tables"] for c in t["columns"])


def test_disable_file_report_removes_stale_report(tmp_path, data):
    export_results(tmp_path, data, ExportOptions(missing_report="file"))
    export_results(tmp_path, data, ExportOptions(missing_report="terminal"), emit=lambda _: None)
    assert not (tmp_path / "missing_values.json").exists()
    assert not (tmp_path / "missing_values.csv").exists()


def test_cli_resume_inherits_and_overrides_without_network(tmp_path, data, monkeypatch, capsys):
    row = make_row(1, issn="0929-8665")
    class OfflineClient:
        mode = "anonymous"
        def __init__(self, *args, **kwargs): pass
        def close(self): pass
        def search(self, row): pytest.fail("Completed task must not request the network")
    monkeypatch.setattr(cli, "LetPubClient", OfflineClient)
    atomic_json(tmp_path / "inputs.json", {"schema_version": 1, "inputs": [row],
                                         "export_options": ExportOptions("zh", True, "file").to_dict()})
    append_event(tmp_path, make_event(row, "matched", "精确匹配", data["journals"][0]))
    assert cli.main(["query", "--resume", str(tmp_path)]) == 0
    assert json.loads((tmp_path / "results.json").read_text())["export_view"]["field_names"] == "zh"
    assert (tmp_path / "missing_values.csv").exists()
    assert "缺失值统计" not in capsys.readouterr().out
    assert cli.main(["query", "--resume", str(tmp_path), "--field-names", "original",
                     "--no-drop-empty-columns", "--missing-report", "terminal"]) == 0
    assert "export_view" not in json.loads((tmp_path / "results.json").read_text())
    assert not (tmp_path / "missing_values.json").exists()
    assert "缺失值统计" in capsys.readouterr().out
    assert json.loads((tmp_path / "inputs.json").read_text())["export_options"] == ExportOptions(missing_report="terminal").to_dict()
    # Old task manifests without options still resume using the original defaults.
    atomic_json(tmp_path / "inputs.json", {"schema_version": 1, "inputs": [row]})
    assert cli.main(["query", "--resume", str(tmp_path)]) == 0


def test_cli_new_task_persists_options_and_validates_configuration(tmp_path, monkeypatch):
    class OfflineClient:
        mode = "anonymous"
        def __init__(self, *args, **kwargs): pass
        def close(self): pass
        def search(self, row): return {"candidates": [], "complete": True}
    monkeypatch.setattr(cli, "LetPubClient", OfflineClient)
    monkeypatch.setattr("sys.stdin", io.StringIO("Unknown journal\n"))
    directory = tmp_path / "new"
    assert cli.main(["query", "--stdin", "--output", str(directory), "--field-names", "zh",
                     "--drop-empty-columns", "--missing-report", "both"]) == 2
    assert json.loads((directory / "inputs.json").read_text())["export_options"] == ExportOptions("zh", True, "both").to_dict()
    assert (directory / "missing_values.json").exists()
    assert cli.main(["query", "--stdin", "--field-names", "invalid"]) == 1
    assert cli.main(["query", "--stdin", "--missing-report", "invalid"]) == 1
    manifest = json.loads((directory / "inputs.json").read_text())
    manifest["export_options"]["drop_empty_columns"] = "false"
    atomic_json(directory / "inputs.json", manifest)
    assert cli.main(["query", "--resume", str(directory)]) == 1


def test_interrupt_still_exports_selected_view_and_report(tmp_path, data):
    class InterruptedClient:
        def search(self, row): raise KeyboardInterrupt
    row = make_row(1, issn="0929-8665")
    messages = []
    code, result = run_batch([row], tmp_path, InterruptedClient(), emit=messages.append,
                             export_options=ExportOptions("zh", True, "both"))
    assert code == 130 and result["queries"][0]["status"] == "pending"
    saved = json.loads((tmp_path / "results.json").read_text())
    assert saved["export_view"]["field_names"] == "zh"
    assert (tmp_path / "missing_values.csv").exists()
    assert any("缺失值统计" in message for message in messages)


def test_chinese_export_keeps_spreadsheet_formula_inert(tmp_path, data):
    data["queries"][0]["journal_name"] = "=1+1"
    export_results(tmp_path, data, ExportOptions("zh", True))
    wb = load_workbook(tmp_path / "results.xlsx")
    sheet = wb["查询记录"]
    labels = [c.value for c in sheet[1]]
    cell = sheet.cell(2, labels.index("期刊全称") + 1)
    assert cell.data_type == "s" and cell.value == "'=1+1"
    wb.close()
    saved = json.loads((tmp_path / "results.json").read_text())
    assert saved["export_view"]["tables"][2]["rows"][0]["期刊全称"] == "=1+1"
