"""Append newly PDF-available bibliographic records, preserving prior six columns."""
import copy
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .client import atomic_json
from .enrichment import EXPORT_LABELS

HEADERS = ['发表时间', 'Cite key', '文献标题', '所属期刊', 'Goal relevance', 'Citations per year']
KEYS = ['year', 'cite_key', 'title', 'journal', 'relevance', 'citations_per_year']


def sync_pdfs(source, snapshot_path, target, header_row=7, sheet_name=None, work_dir=None):
    source, target = Path(source).resolve(), Path(target).resolve()
    if target == source or target.exists() or target.suffix.lower() != '.xlsx':
        raise ValueError('输出必须为尚不存在的新xlsx文件')
    snapshot = json.loads(Path(snapshot_path).read_text(encoding='utf-8'))
    papers = snapshot.get('all_papers', snapshot.get('papers'))
    if not isinstance(papers, list):
        raise ValueError('PDF快照须包含papers或all_papers列表')
    known = set()
    for p in papers:
        if not isinstance(p, dict) or not p.get('cite_key') or p['cite_key'] in known:
            raise ValueError('PDF快照Cite key为空或重复')
        known.add(p['cite_key'])
        if p.get('pdf') not in (True, False, '✓', 'X', '…'):
            raise ValueError('每篇文献须明确提供PDF状态')
    wb = load_workbook(source)
    try:
        sh = wb[sheet_name] if sheet_name else wb.worksheets[0]
        headers = [c.value for c in sh[header_row]]
        if headers[:6] != HEADERS or (len(headers) > 6 and headers[6:] != EXPORT_LABELS):
            raise ValueError('仅支持原六字段或本工具完整补充的文献表')
        table = next((t for t in sh.tables.values() if range_boundaries(t.ref)[:2] == (1, header_row)), None)
        if not table or table.totalsRowCount:
            raise ValueError('须有无总计行的原生文献表')
        end = range_boundaries(table.ref)[3]
        old = [[sh.cell(r,c).value for c in range(1,7)] for r in range(header_row+1,end+1)]
        keys = [r[1] for r in old]
        if any(not k for k in keys) or len(set(keys)) != len(keys):
            raise ValueError('原表Cite key为空或重复')
        additions = [p for p in papers if p['pdf'] in (True, '✓') and p['cite_key'] not in set(keys)]
        for p in additions:
            if any(k not in p for k in KEYS) or not p['title'] or not p['journal']:
                raise ValueError('新增文献缺少原六字段')
            for k in ('year', 'relevance', 'citations_per_year'):
                if p[k] is not None and (not isinstance(p[k], (int,float)) or isinstance(p[k], bool)):
                    raise ValueError('年份及指标须为数值或null')
        # Reusable output is the six-column source for a subsequent enrich run.
        if len(headers)>6:
            for merged in list(sh.merged_cells.ranges):
                if merged.min_col>6: sh.unmerge_cells(str(merged))
            sh.delete_cols(7,sh.max_column-6)
            table.tableColumns=table.tableColumns[:6]
        for i,p in enumerate(additions,end+1):
            for c,k in enumerate(KEYS,1):
                cell=sh.cell(i,c); previous=sh.cell(end,c)
                cell._style=copy.copy(previous._style)
                cell.value=p[k]
                if isinstance(p[k],str): cell.data_type='s'
            sh.row_dimensions[i].height=max(70,sh.row_dimensions[end].height or 15)
        last=end+len(additions);table.ref=f'A{header_row}:F{last}'
        if isinstance(sh["A4"].value,str) and sh["A4"].value.startswith("截至 "):
            sh["A4"] = f"PDF增量：保留原{len(old)}篇，新增{len(additions)}篇，共{last-header_row}篇；原六列数据与旧记录顺序保留。"
        table.autoFilter.ref=table.ref
        assert old==[[sh.cell(r,c).value for c in range(1,7)] for r in range(header_row+1,end+1)]
        report={'old_count':len(old),'new_count':len(additions),'total_count':last-header_row,
                'new_keys':[p['cite_key'] for p in additions],
                'previous_keys_absent_from_snapshot':sorted(set(keys)-known),
                'prior_rows_preserved':True,'source':str(source),'output':str(target)}
        target.parent.mkdir(parents=True,exist_ok=True)
        temp=target.with_name('.'+target.stem+'.tmp.xlsx')
        try: wb.save(temp);temp.replace(target)
        finally: temp.unlink(missing_ok=True)
        if work_dir: atomic_json(Path(work_dir)/'pdf_sync_report.json',report)
        return report
    finally:
        wb.close()
