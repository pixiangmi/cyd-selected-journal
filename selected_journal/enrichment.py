"""Enrich a literature XLSX using deduplicated, traceable LetPub records."""
import copy
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import TableColumn

from .client import atomic_json
from .inputs import name_key, normalize_issn, normalize_name
from .storage import HEADER_LABELS, SUMMARY_HEADERS, projections, write_csv, value_for_table

EXTRA = {
    "match_status": "期刊匹配状态", "match_message": "匹配说明",
    "verification_source": "刊名核查来源",
    "jcr_year": "JCR分区_年份",
    "journal_partition_year": "期刊分区表（中科院）_年份",
    "xinrui_year": "新锐分区_年份",
}
PRIMARY = ["journal_name", "abbreviation", "issn", "eissn", "impact_factor",
           "citescore", "oa", "review_time", "jcr_subjects",
           "journal_partition_subjects", "xinrui_subjects", "match_status",
           "match_message", "verification_source", "source_url", "fetched_at"]
EXPORT_KEYS = PRIMARY + [k for k in SUMMARY_HEADERS if k not in PRIMARY] + [
    k for k in EXTRA if k not in PRIMARY]
EXPORT_LABELS = [EXTRA.get(k, HEADER_LABELS.get(k, k)) for k in EXPORT_KEYS]
PREPRINTS = {"biorxiv", "medrxiv", "arxiv", "osf preprints"}
UNKNOWN = {"", "unknown journal", "unknown", "未知期刊", "期刊未知"}


def load_mappings(path):
    data = json.loads(Path(path).read_text(encoding="utf-8")) if path else {}
    aliases = {}
    for original, mapping in data.get("verified_aliases", {}).items():
        key = name_key(original)
        if key in aliases:
            raise ValueError("核查映射包含重复规范刊名")
        aliases[key] = mapping
    overrides = data.get("paper_overrides", {})
    for mapping in [*aliases.values(), *overrides.values()]:
        if mapping.get("record_type", "journal") not in ("journal", "preprint", "unknown"):
            raise ValueError("record_type 必须为 journal、preprint 或 unknown")
        if not (mapping.get("source") or mapping.get("sources")):
            raise ValueError("核查映射必须提供 source 或 sources")
        if mapping.get("issn"):
            normalize_issn(mapping["issn"])
    return aliases, overrides


def extract(path, journal_column="所属期刊", key_column=None, header_row=None,
            sheet_name=None, mappings=None):
    path = Path(path).resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("文献回填仅支持 .xlsx")
    wb = load_workbook(path)
    sheet = wb[sheet_name] if sheet_name else wb.worksheets[0]
    try:
        if header_row is None:
            found = [r for r in range(1, min(sheet.max_row, 50) + 1)
                     if any(c.value == journal_column for c in sheet[r])]
            if len(found) != 1:
                raise ValueError("不能唯一定位期刊表头，请指定 --header-row")
            header_row = found[0]
        headers = [c.value for c in sheet[header_row]]
        if headers.count(journal_column) != 1:
            raise ValueError("期刊列不存在或重复")
        jcol = headers.index(journal_column) + 1
        if key_column is None and "Cite key" in headers:
            key_column = "Cite key"
        if key_column and headers.count(key_column) != 1:
            raise ValueError("记录标识列不存在或重复")
        kcol = headers.index(key_column) + 1 if key_column else None
        aliases, overrides = load_mappings(mappings)
        table = next((t for t in sheet.tables.values()
                      if range_boundaries(t.ref)[1] == header_row
                      and range_boundaries(t.ref)[0] <= jcol <= range_boundaries(t.ref)[2]), None)
        end_row = range_boundaries(table.ref)[3] if table else sheet.max_row
        if table and table.totalsRowCount:
            raise ValueError("含总计行的文献表暂不支持回填")
        rows, seen, queries = [], set(), {}
        for row in range(header_row + 1, end_row + 1):
            values = [sheet.cell(row, c).value for c in range(1, len(headers) + 1)]
            if all(v is None for v in values):
                continue
            if sheet.cell(row, jcol).data_type == "f":
                raise ValueError("期刊名称不能是待计算公式")
            key_value = sheet.cell(row, kcol).value if kcol else None
            key = str(key_value if key_value is not None else "").strip() if kcol else f"row:{row}"
            if not key or key in seen:
                raise ValueError("记录标识为空或重复")
            seen.add(key)
            original = normalize_name(sheet.cell(row, jcol).value)
            mapping = overrides.get(key, aliases.get(name_key(original), {}))
            name = normalize_name(mapping.get("standard_name", mapping.get("name", original)))
            kind = mapping.get("record_type") or ("preprint" if name_key(name) in PREPRINTS
                    else "unknown" if name_key(name) in UNKNOWN else "journal")
            issn = normalize_issn(mapping["issn"]) if mapping.get("issn") else ""
            evidence = mapping.get("sources") or ([mapping["source"]] if mapping.get("source") else [])
            spec = {"journal_name": "" if issn else name, "issn": issn}
            query_key = "issn:" + issn if issn else "name:" + name_key(name)
            if kind == "journal":
                queries.setdefault(query_key, spec)
            rows.append({"excel_row": row, "key": key, "original_name": original,
                         "normalized_name": name_key(original), "query_key": query_key,
                         "standard_name": name, "record_type": kind,
                         "verification_sources": evidence, "verification_note": mapping.get("note", "")})
        return wb, sheet, {"source_file": str(path),
                "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "sheet": sheet.title, "header_row": header_row, "headers": headers,
                "table": table.name if table else None, "rows": rows,
                "queries": list(queries.values())}
    except Exception:
        wb.close()
        raise


def prepare(manifest, directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    atomic_json(directory / "literature_inputs.json", manifest)
    write_csv(directory / "journals.csv", ["journal_name", "issn"],
              [[q["journal_name"], q["issn"]] for q in manifest["queries"]])
    grouped = {}
    for row in manifest["rows"]:
        group = grouped.setdefault(row["original_name"], {"name": row["original_name"],
                                    "keys": [], "normalized": row["normalized_name"], "queries": set()})
        group["keys"].append(row["key"])
        group["queries"].add(row["query_key"])
    write_csv(directory / "期刊去重统计.csv", ["原始期刊名称", "规范去重名称", "查询名称或ISSN", "文献篇数", "文献标识"],
              [[g["name"], g["normalized"], "; ".join(sorted(g["queries"])), len(g["keys"]), "; ".join(g["keys"])] for g in grouped.values()])
    return {"papers": len(manifest["rows"]), "original_names": len(grouped),
            "normalized_names": len({r["normalized_name"] for r in manifest["rows"]}),
            "queries": len(manifest["queries"])}


def read_results(paths):
    journals, queries, origins = {}, {}, {}
    ranks = {"matched": 4, "partial": 3, "needs_confirmation": 2, "not_found": 1}
    for path in paths:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        if data.get("schema_version") != 1:
            raise ValueError("期刊结果 schema_version 不受支持")
        for j in data["journals"]:
            old = journals.get(j["journal_id"])
            score = lambda x: sum(f["status"] == "ok" for f in x["fields"].values()) + sum(
                p["status"] == "ok" for p in x["classifications"])
            if old is None or score(j) >= score(old):
                journals[j["journal_id"]] = j
                origins[j["journal_id"]] = str(Path(path).resolve())
        for q in data["queries"]:
            key = "issn:" + normalize_issn(q["issn"]) if q.get("issn") else "name:" + name_key(q["journal_name"])
            if key not in queries or ranks.get(q["status"], 0) >= ranks.get(queries[key]["status"], 0):
                queries[key] = q
    summary = projections({"journals": list(journals.values()), "queries": [], "candidates": []})[0][2]
    return journals, queries, {r["journal_id"]: r for r in summary}, origins


def join_rows(manifest, result_paths):
    journals, queries, summaries, origins = read_results(result_paths)
    outputs, mappings = [], []
    for row in manifest["rows"]:
        status, message, journal_id = "unknown", "未确认期刊名称", None
        record = None
        if row["record_type"] == "preprint":
            status, message = "preprint", "预印本平台；期刊指标不适用"
        elif row["record_type"] == "journal":
            q = queries.get(row["query_key"])
            if q:
                status, message, journal_id = q["status"], q["message"], q.get("journal_id")
            else:
                status, message = "pending", "缺少对应期刊查询结果"
            if journal_id and status in ("matched", "partial"):
                record = journals.get(journal_id)
            if record is None:
                # Another spelling or ISSN query may have resolved the same exact identity.
                issn = row["query_key"][5:] if row["query_key"].startswith("issn:") else None
                candidates = [j for j in journals.values() if (
                    issn and issn in (j["fields"]["issn"]["value"], j["fields"]["eissn"]["value"])) or (
                    not issn and name_key(j["fields"]["journal_name"]["value"]) == name_key(row["standard_name"]))]
                if len(candidates) == 1:
                    record = candidates[0]
                    journal_id = record["journal_id"]
                    broken = any(f["status"] == "parse_error" for f in record["fields"].values()) or any(
                        p["status"] == "parse_error" for p in record["classifications"])
                    status, message = ("partial" if broken else "matched"), "已核实身份与结果中的期刊一致"
        if record and status == "partial":
            broken = any(f["status"] == "parse_error" for f in record["fields"].values()) or any(
                p["status"] == "parse_error" for p in record["classifications"])
            if not broken:
                status, message = "matched", "已核实同一期刊；合并采用完整复查记录"
        values = dict(summaries.get(journal_id, {})) if record else {}
        if record:
            for partition in record["classifications"]:
                values[partition["scheme"] + "_year"] = partition.get("year")
        else:
            if row["record_type"] == "preprint" or row["verification_sources"]:
                values["journal_name"] = row["standard_name"]
            values["journal_name_status"] = "ok" if values.get("journal_name") else "not_provided"
            for key in SUMMARY_HEADERS:
                if key.endswith("_status") and key != "journal_name_status":
                    values[key] = "not_applicable" if status == "preprint" else "not_queried"
        if row["verification_note"]:
            message += "；" + row["verification_note"]
        values.update({"match_status": status, "match_message": message,
                       "verification_source": "; ".join(row["verification_sources"]) or None})
        outputs.append([None if values.get(k) == "" else value_for_table(values.get(k)) for k in EXPORT_KEYS])
        mappings.append({**row, "match_status": status, "match_message": message,
                         "journal_id": journal_id, "result_source": origins.get(journal_id)})
    return outputs, mappings


def enrich(args, query_main):
    source, target = args.input.resolve(), args.output.resolve()
    if source == target:
        raise ValueError("输出必须是新文件，不能覆盖原文献表")
    if target.exists():
        raise ValueError("输出文件已存在，请使用新的输出路径")
    if target.suffix.lower() != ".xlsx":
        raise ValueError("输出必须为 .xlsx")
    wb, sheet, manifest = extract(source, args.journal_column, args.key_column,
                                 args.header_row, args.sheet, args.mappings)
    support = args.work_dir.resolve() if args.work_dir else target.parent / (target.stem + "_资料")
    try:
        if any(label in manifest["headers"] for label in EXPORT_LABELS):
            raise ValueError("表内已有期刊补充字段，请使用未补充的原始文献表")
        statistics = prepare(manifest, support)
        result_paths = list(args.results or [])
        if not result_paths and manifest["queries"]:
            directory = support / "query"
            command = ["query", "--field-names", "zh", "--no-drop-empty-columns", "--missing-report", "both"]
            if (directory / "inputs.json").exists():
                previous = json.loads((directory / "inputs.json").read_text())
                old = [(r["journal_name"], r["issn"]) for r in previous["inputs"]]
                new = [(q["journal_name"], q["issn"]) for q in manifest["queries"]]
                if old != new:
                    raise ValueError("已有查询任务与当前去重名单不同，请使用新的 --work-dir")
                command += ["--resume", str(directory)]
            else:
                command += ["--input", str(support / "journals.csv"), "--output", str(directory)]
            if args.cookie_file:
                command += ["--cookie-file", str(args.cookie_file)]
            if args.refresh:
                command += ["--refresh"]
            if getattr(args, "request_interval", None) is not None:
                command += ["--request-interval", str(args.request_interval)]
            for identity_path in getattr(args, "identity_results", None) or []:
                command += ["--identity-results", str(identity_path)]
            code = query_main(command)
            if code == 1:
                raise ValueError("期刊查询配置失败")
            result_paths = [directory / "results.json"]
        outputs, mappings = join_rows(manifest, result_paths)
        start = len(manifest["headers"]) + 1
        original_rows = {r["excel_row"]: [sheet.cell(r["excel_row"], c).value for c in range(1, start)]
                         for r in manifest["rows"]}
        header_style = copy.copy(sheet.cell(manifest["header_row"], start - 1)._style)
        for col, label in enumerate(EXPORT_LABELS, start):
            cell = sheet.cell(manifest["header_row"], col, label)
            cell._style = copy.copy(header_style)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            key = EXPORT_KEYS[col - start]
            width = 58 if key.endswith("_raw") or key in ("match_message", "verification_source", "source_url") else (
                48 if key.endswith("_subjects") else 34 if key in ("journal_name", "review_time") else 24)
            sheet.column_dimensions[get_column_letter(col)].width = width
        last_col = start + len(EXPORT_KEYS) - 1
        for row, values in zip(manifest["rows"], outputs):
            rownum = row["excel_row"]
            height = sheet.row_dimensions[rownum].height or 15
            for col, value in enumerate(values, start):
                cell = sheet.cell(rownum, col, value)
                cell.font = copy.copy(sheet.cell(rownum, start - 1).font)
                numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
                cell.alignment = Alignment(horizontal="right" if numeric else "left",
                                           vertical="top", wrap_text=not numeric)
                key = EXPORT_KEYS[col - start]
                if key in ("impact_factor", "citescore"):
                    cell.number_format = "0.000" if key == "impact_factor" else "0.0"
                elif key.endswith("_year"):
                    cell.number_format = "0"
                if value is not None and not numeric:
                    width = sheet.column_dimensions[get_column_letter(col)].width
                    chunks = str(value).splitlines() or [""]
                    units = lambda s: sum(2 if ord(c) > 255 else 1 for c in s)
                    lines = sum(max(1, math.ceil(units(s) / max(1, width - 3))) for s in chunks)
                    height = max(height, lines * 15 + 5)
            sheet.row_dimensions[rownum].height = min(409, height)
        sheet.row_dimensions[manifest["header_row"]].height = max(
            sheet.row_dimensions[manifest["header_row"]].height or 15, 60)
        last_row = max((r["excel_row"] for r in manifest["rows"]), default=manifest["header_row"])
        if manifest["table"]:
            table = sheet.tables[manifest["table"]]
            left, top, _, bottom = range_boundaries(table.ref)
            table.ref = f"{get_column_letter(left)}{top}:{get_column_letter(last_col)}{bottom}"
            if table.autoFilter:
                table.autoFilter.ref = table.ref
            table.tableColumns.extend(TableColumn(id=c - left + 1, name=label)
                                      for c, label in enumerate(EXPORT_LABELS, start))
        else:
            sheet.auto_filter.ref = f"A{manifest['header_row']}:{get_column_letter(last_col)}{last_row}"
        if not sheet.freeze_panes:
            sheet.freeze_panes = f"A{manifest['header_row'] + 1}"
        for rownum, values in original_rows.items():
            if values != [sheet.cell(rownum, c).value for c in range(1, start)]:
                raise ValueError("原列值发生变化")
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name("." + target.stem + ".tmp.xlsx")
        try:
            wb.save(temporary)
            check = load_workbook(temporary, read_only=True, data_only=False)
            try:
                out_sheet = check[sheet.title]
                for rownum, values in enumerate(out_sheet.iter_rows(
                        min_row=manifest["header_row"] + 1, max_col=start - 1, values_only=True),
                        manifest["header_row"] + 1):
                    if rownum in original_rows and original_rows[rownum] != list(values):
                        raise ValueError("导出后原列值发生变化")
            finally:
                check.close()
            temporary.replace(target)
        finally:
            temporary.unlink(missing_ok=True)
        atomic_json(support / "row_mapping.json", mappings)
        atomic_json(support / "enrichment_values.json", {"keys": EXPORT_KEYS, "labels": EXPORT_LABELS,
                    "values": outputs, "excel_rows": [r["excel_row"] for r in manifest["rows"]]})
        metrics = {}
        for key in ("impact_factor", "citescore", "jcr_subjects", "journal_partition_subjects", "xinrui_subjects"):
            col = EXPORT_KEYS.index(key)
            metrics[key] = {"missing_papers": sum(v[col] is None or v[col] == "" for v in outputs),
                           "total_papers": len(outputs)}
        report = {**statistics, "matched_journals": len({m["journal_id"] for m in mappings if m["journal_id"]}),
                  "matched_papers": sum(m["journal_id"] is not None for m in mappings),
                  "statuses": dict(Counter(m["match_status"] for m in mappings)),
                  "metrics": metrics, "output": str(target), "source_sha256": manifest["source_sha256"]}
        atomic_json(support / "enrichment_report.json", report)
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if all(m["match_status"] in ("matched", "preprint") for m in mappings) else 2
    finally:
        wb.close()

