"""Input validation; invalid rows remain visible in the query report."""
import csv
import re
import unicodedata
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook


def normalize_name(value):
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split())


def name_key(value):
    return normalize_name(value).casefold()


def normalize_issn(value):
    text = normalize_name(value).upper()
    digits = re.sub(r"[\s\-‐‑–—]", "", text)
    if not re.fullmatch(r"\d{7}[\dX]", digits):
        raise ValueError("ISSN 格式应为 XXXX-XXXX（末位允许 X）")
    check = (11 - sum(int(digits[i]) * (8 - i) for i in range(7)) % 11) % 11
    if digits[-1] != ("X" if check == 10 else str(check)):
        raise ValueError("ISSN 校验位不正确")
    return digits[:4] + "-" + digits[4:]


def make_row(index, name="", issn="", raw=None):
    row = {"input_id": index, "journal_name": normalize_name(name),
           "issn": normalize_name(issn), "raw": raw, "error": None}
    if not row["journal_name"] and not row["issn"]:
        return None
    if row["issn"]:
        try:
            row["issn"] = normalize_issn(row["issn"])
        except ValueError as error:
            row["error"] = str(error)
    return row


def read_lines(lines):
    rows = []
    for index, raw in enumerate(lines, 1):
        text = normalize_name(raw)
        if not text:
            continue
        # ISSN-shaped input with an invalid check digit must not become a name.
        is_issn = bool(re.fullmatch(r"[\dXx\s\-‐‑–—]{8,}", text))
        row = make_row(index, "" if is_issn else text, text if is_issn else "", text)
        rows.append(row)
    return rows


def _table_rows(headers, values):
    columns = [normalize_name(h).lower() for h in headers]
    if not ({"journal_name", "issn"} & set(columns)):
        raise ValueError("表格必须有 journal_name 或 issn 列")
    if len(columns) != len(set(columns)):
        raise ValueError("表格列名不能重复")
    rows = []
    for index, values_row in enumerate(values, 2):
        data = dict(zip(columns, values_row))
        row = make_row(index, data.get("journal_name"), data.get("issn"))
        if row:
            rows.append(row)
    return rows


def read_file(filename):
    path = Path(filename)
    if path.suffix.lower() == ".txt":
        return read_lines(path.read_text(encoding="utf-8-sig").splitlines())
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            return _table_rows(next(reader, []), reader)
    if path.suffix.lower() == ".xlsx":
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except (BadZipFile, KeyError, ValueError) as error:
            raise ValueError("无法读取 XLSX，请检查文件是否损坏或格式错误") from error
        try:
            values = workbook.worksheets[0].iter_rows(values_only=True)
            return _table_rows(next(values, []), values)
        finally:
            workbook.close()
    raise ValueError("输入文件仅支持 TXT、CSV、XLSX")
