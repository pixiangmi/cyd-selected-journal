"""Append-only checkpoints and consistent Excel/CSV/JSON projections."""
import copy
import csv
import json
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .client import atomic_json

FIELD_NAMES = ("journal_name", "abbreviation", "issn", "eissn", "impact_factor", "citescore", "oa", "review_time")
SUMMARY_HEADERS = ["journal_id", "input_ids", "source_url", "fetched_at"] + [
    f"{name}{suffix}" for name in FIELD_NAMES for suffix in ("", "_status", "_raw", "_version", "_year")]
SUMMARY_HEADERS += [f"{scheme}_{part}" for scheme in ("jcr", "journal_partition", "xinrui")
                    for part in ("source_label", "version", "status", "subjects")]
PARTITION_HEADERS = ["journal_id", "journal_name", "scheme", "source_label", "version", "year", "status",
                     "level", "subject", "quartile", "basis", "source_url", "fetched_at"]
QUERY_HEADERS = ["input_id", "journal_name", "issn", "raw", "status", "journal_id", "message"]
CANDIDATE_HEADERS = ["input_id", "journal_id", "journal_name", "issn", "url", "reason"]


def append_event(directory, event):
    with (Path(directory) / "checkpoint.jsonl").open("a", encoding="utf-8") as stream:
        stream.write(json.dumps(event, ensure_ascii=False) + "\n")
        stream.flush()
        os.fsync(stream.fileno())


def read_events(directory):
    path = Path(directory) / "checkpoint.jsonl"
    if not path.exists():
        return {}
    data = path.read_bytes()
    events = {}
    offset = 0
    for index, line in enumerate(data.splitlines(keepends=True)):
        try:
            event = json.loads(line)
            events[event["query"]["input_id"]] = event
        except (ValueError, KeyError, TypeError):
            if index != len(data.splitlines()) - 1:
                raise ValueError("检查点中间记录损坏，无法安全恢复")
            # Only discard a torn trailing write, retaining all prior events.
            with path.open("r+b") as stream:
                stream.truncate(offset)
            break
        offset += len(line)
    else:
        if data and not data.endswith(b"\n"):
            with path.open("ab") as stream:
                stream.write(b"\n")
    return events


def result_data(inputs, events):
    queries = []
    journals = {}
    candidates = []
    for row in inputs:
        event = events.get(row["input_id"])
        if event is None:
            queries.append({**row, "status": "pending", "journal_id": None, "message": "尚未处理"})
            continue
        queries.append(event["query"])
        candidates.extend(event.get("candidates", []))
        record = event.get("journal")
        if record:
            key = record["journal_id"]
            if key not in journals:
                journals[key] = copy.deepcopy(record)
                journals[key]["input_ids"] = []
            journals[key]["input_ids"].append(row["input_id"])
    return {"schema_version": 1, "journals": list(journals.values()), "queries": queries, "candidates": candidates}


def projections(data):
    summary, partitions = [], []
    for journal in data["journals"]:
        row = {key: journal[key] for key in ("journal_id", "input_ids", "source_url", "fetched_at")}
        for name, f in journal["fields"].items():
            row[name] = f["value"]
            for part in ("status", "raw", "version", "year"):
                row[f"{name}_{part}"] = f[part]
        for partition in journal["classifications"]:
            scheme = partition["scheme"]
            for part in ("source_label", "version", "status"):
                row[f"{scheme}_{part}"] = partition[part]
            row[f"{scheme}_subjects"] = "; ".join(
                f"{e['basis'] or e['level']}: {e['subject']} {e['quartile'] or ''}" for e in partition["entries"])
            for entry in partition["entries"] or [{}]:
                partitions.append({**{k: partition[k] for k in ("scheme", "source_label", "version", "year", "status")},
                                   **entry, "journal_id": journal["journal_id"],
                                   "journal_name": journal["fields"]["journal_name"]["value"],
                                   "source_url": journal["source_url"], "fetched_at": journal["fetched_at"]})
        summary.append(row)
    return [("期刊汇总", SUMMARY_HEADERS, summary), ("分区明细", PARTITION_HEADERS, partitions),
            ("查询记录", QUERY_HEADERS, data["queries"]), ("候选清单", CANDIDATE_HEADERS, data["candidates"])]


def value_for_table(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        # Preserve raw originals in JSON while making spreadsheet cells inert.
        value = "".join(c for c in value if ord(c) >= 32 or c in "\n\r\t")
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
    return value


def export_results(directory, data):
    directory = Path(directory)
    atomic_json(directory / "results.json", data)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in projections(data):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([value_for_table(row.get(h)) for h in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24557A")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(50, max(14, max(len(str(c.value or "")) for c in col) + 2))
        handle, temporary = tempfile.mkstemp(prefix=".tmp-", dir=directory)
        try:
            with os.fdopen(handle, "w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                writer.writerows([value_for_table(row.get(h)) for h in headers] for row in rows)
            os.replace(temporary, directory / f"{title}.csv")
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
    handle, temporary = tempfile.mkstemp(prefix=".tmp-", suffix=".xlsx", dir=directory)
    os.close(handle)
    try:
        workbook.save(temporary)
        os.replace(temporary, directory / "results.xlsx")
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)
